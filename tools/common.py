"""Sheets every Northwall model shares: Cover, Valuation frame (CAPM, WACC, target, rating,
CFA cross-checks), Sensitivity, Ratios, Comps, Checks, Log."""
from openpyxl.styles import Font, Alignment
from engine import (PERIODS, PCOL, FORECAST, FMT, F_IN, F_CALC, F_LINK, F_B, F_NOTE, FILL_KEY,
                    RULE, Font as _F, FONT, NAVY, is_formula)

YEARS = FORECAST  # 2027E..2030E explicit forecast for DCF / RI


def valuation(m, arch):
    c = m.cfg
    ws = m.sheet('Valuation', [46, 16, 70])
    m.title(ws, f"{c['company']['name']} — Valuation and price target",
            'Target methods first; CFA Level II cross-checks (DCF, reverse DCF, residual income, justified P/E) follow.')
    S, M = m.S, m.M
    r = 4
    r = m.block(ws, r, [
        ('#', 'Market data'),
        ('price', f"Share price ({c['company']['price_ccy']})", f"={S('price')}", 'ps', c['scalars']['price']['source']),
        ('price_model', f"Share price in model currency ({c['company']['model_ccy']})", f"={S('price')}*{S('fx_price')}", 'ps', 'Converted when the listing currency differs from the reporting currency'),
        ('shares', 'Basic shares outstanding (M)', f"={S('shares')}", 'n1', ''),
        ('mcap', 'Market capitalization (M)', "=B6*B7".replace('B6', 'B6'), 'm0', ''),
        ('netdebt_now', 'Net debt (cash) at last balance sheet (M)', f"={S('netdebt_now')}", 'm0', 'Negative = net cash'),
        ('ev', 'Enterprise value (M)', None, 'm0', 'Market cap + net debt'),
    ])
    # fix references now that cell addresses exist
    ws[m.ref('mcap').split('!')[1].replace('$', '')] = f"={m.ref('price_model')}*{m.ref('shares')}"
    ws[m.ref('ev').split('!')[1].replace('$', '')] = f"={m.ref('mcap')}+{m.ref('netdebt_now')}"
    r = m.block(ws, r + 1, [
        ('#', 'Required return on equity (CAPM) and WACC — CFA L2 Equity: Return Concepts; Corporate Issuers: Cost of Capital'),
        ('rf', 'Risk-free rate (US 10-yr)', f"={S('rf')}", 'pct', 'House deck'),
        ('beta', 'Equity beta', f"={S('beta')}", 'n2', c['scalars']['beta']['source']),
        ('erp', 'Equity risk premium', f"={S('erp')}", 'pct', 'House deck'),
        ('ke', 'Cost of equity: rf + beta x ERP', None, 'pct', 'CAPM'),
        ('kd', 'Pre-tax cost of debt', f"={S('kd')}", 'pct', c['scalars']['kd']['source']),
        ('tax', 'Marginal tax rate', f"={S('tax_rate')}", 'pct', ''),
        ('debt_w', 'Debt weight: D / (D + E), market values', None, 'pct', 'Gross debt at market cap'),
        ('wacc', 'WACC', None, 'pct', 'ke x E/V + kd x (1 - t) x D/V'),
    ])
    ws[m.ref('ke').split('!')[1].replace('$', '')] = f"={m.ref('rf')}+{m.ref('beta')}*{m.ref('erp')}"
    ws[m.ref('debt_w').split('!')[1].replace('$', '')] = f"={S('gross_debt')}/({S('gross_debt')}+{m.ref('mcap')})"
    ws[m.ref('wacc').split('!')[1].replace('$', '')] = (f"={m.ref('ke')}*(1-{m.ref('debt_w')})+{m.ref('kd')}*(1-{m.ref('tax')})*{m.ref('debt_w')}")

    # ---- archetype-specific target methods
    r = arch.target_methods(m, ws, r + 1)
    r = m.block(ws, r, [
        ('#', 'Price target and rating'),
        ('w1', 'Weight on method 1', f"={S('w1')}", 'pct', ''),
        ('w2', 'Weight on method 2', f"=1-{S('w1')}", 'pct', ''),
        ('blend', f"Blended value per share ({c['company']['model_ccy']})", None, 'ps', ''),
        ('blend_px', f"Blended value per share ({c['company']['price_ccy']})", None, 'ps', 'Converted to listing currency'),
        ('target', f"12-month price target ({c['company']['price_ccy']}), rounded", None, 'ps', 'Rounded to the house increment'),
        ('dps_fwd', 'Next-12-month dividend per share (listing currency)', f"={S('dps_fwd')}", 'ps', c['scalars']['dps_fwd']['source']),
        ('ret_px', 'Implied price return', None, 'pct', ''),
        ('ret_div', 'Dividend yield', None, 'pct', ''),
        ('ret_all', 'Implied all-in return', None, 'pct', 'Rating bands: Buy > 15%, Hold 0–15%, Sell < 0%'),
        ('rating', 'Rating', None, 'txt', ''),
    ])
    cell = lambda k: ws[m.ref(k).split('!')[1].replace('$', '')]
    cell('blend').value = f"={m.ref('w1')}*{m.ref('v1')}+{m.ref('w2')}*{m.ref('v2')}"
    cell('blend_px').value = f"={m.ref('blend')}/{S('fx_price')}"
    cell('target').value = f"=MROUND({m.ref('blend_px')},{S('round_to')})"
    cell('target').fill = FILL_KEY
    cell('ret_px').value = f"={m.ref('target')}/{m.ref('price')}-1"
    cell('ret_div').value = f"={m.ref('dps_fwd')}/{m.ref('price')}"
    cell('ret_all').value = f"={m.ref('ret_px')}+{m.ref('ret_div')}"
    cell('rating').value = f'=IF({m.ref("ret_all")}>{S("buy_band")},"Buy",IF({m.ref("ret_all")}>=0,"Hold","Sell"))'
    cell('rating').fill = FILL_KEY

    # ---- CFA cross-check: two-stage FCFF DCF
    r += 1
    m.section(ws, r, 'Cross-check 1: two-stage FCFF DCF — CFA L2 Equity: Free Cash Flow Valuation', 3); r += 1
    ws.cell(row=r, column=1, value='Valuation date: Dec 31, 2026. FCFF discounted at WACC, end-year; Gordon growth terminal value on 2030E FCFF.').font = F_NOTE
    r += 1
    hdr = r
    ws.cell(row=hdr, column=1, value='Year').font = F_B
    for i, y in enumerate(YEARS):
        col = 2 + i
        ws.cell(row=hdr, column=col, value=y).font = F_B
        ws.cell(row=hdr + 1, column=1, value=f"FCFF ({c['company']['model_ccy']} M)")
        x = ws.cell(row=hdr + 1, column=col, value=f"={M('fcff', y)}"); x.font = F_LINK; x.number_format = FMT['m0']
        ws.cell(row=hdr + 2, column=1, value='Discount factor')
        x = ws.cell(row=hdr + 2, column=col, value=f"=1/(1+{m.ref('wacc')})^{i + 1}"); x.number_format = '0.000'
        ws.cell(row=hdr + 3, column=1, value='Present value')
        x = ws.cell(row=hdr + 3, column=col, value=f"={ws.cell(row=hdr + 1, column=col).coordinate}*{ws.cell(row=hdr + 2, column=col).coordinate}"); x.number_format = FMT['m0']
    for rr in range(hdr, hdr + 4):
        ws.column_dimensions['D'].width = 12; ws.column_dimensions['E'].width = 12
    r = hdr + 5
    pv_sum = f"SUM(B{hdr + 3}:E{hdr + 3})"
    r = m.block(ws, r, [
        ('g', 'Terminal growth, nominal', f"={S('g_term')}", 'pct', c['scalars']['g_term']['source']),
        ('pv_explicit', 'PV of 2027E–2030E FCFF', f"={pv_sum}", 'm0', ''),
        ('tv', 'Terminal value at 2030', f"=E{hdr + 1}*(1+B{r})/({m.ref('wacc')}-B{r})", 'm0', 'FCFF 2030 x (1 + g) / (WACC - g)'),
        ('pv_tv', 'PV of terminal value', None, 'm0', ''),
        ('ev_dcf', 'Enterprise value (DCF)', None, 'm0', ''),
        ('nd_ye26', 'Less: net debt at Dec 31, 2026E', f"={M('netdebt', '2026E')}", 'm0', 'From Model'),
        ('inv', 'Plus: investments not in FCFF', f"={S('investments')}", 'm0', ''),
        ('dcf_ps', f"DCF value per share ({c['company']['model_ccy']})", None, 'ps', 'Cross-check only; not in the target weights'),
        ('dcf_vs', 'DCF value vs. share price', None, 'pct', ''),
        ('tv_share', 'Terminal value share of EV', None, 'pct', 'A high share means the answer rests on long-run assumptions'),
    ])
    cell('tv').value = f"=E{hdr + 1}*(1+{m.ref('g')})/({m.ref('wacc')}-{m.ref('g')})"
    cell('pv_tv').value = f"={m.ref('tv')}*E{hdr + 2}"
    cell('ev_dcf').value = f"={m.ref('pv_explicit')}+{m.ref('pv_tv')}"
    cell('dcf_ps').value = f"=({m.ref('ev_dcf')}-{m.ref('nd_ye26')}+{m.ref('inv')})/{m.ref('shares')}"
    cell('dcf_vs').value = f"={m.ref('dcf_ps')}/{m.ref('price_model')}-1"
    cell('tv_share').value = f"={m.ref('pv_tv')}/{m.ref('ev_dcf')}"
    m.dcf_rows = hdr  # for sensitivity

    # ---- reverse DCF: market-implied growth
    r = m.block(ws, r + 1, [
        ('#', 'Cross-check 2: market-implied growth (reverse Gordon) — CFA L2 Equity: DDM / FCF valuation'),
        ('cf1', f"Next-year cash flow per share, 2027E ({c['company']['model_ccy']})", f"={M('cfps', '2027E')}", 'ps', 'Free cash flow to equity proxy'),
        ('g_impl', 'Growth rate implied by today\'s price: ke - CF1 / P', None, 'pct', 'Compare with the growth the model actually forecasts'),
    ])
    cell('g_impl').value = f"={m.ref('ke')}-{m.ref('cf1')}/{m.ref('price_model')}"

    # ---- residual income and justified P/E (where earnings and book value are modelled)
    if c['company'].get('has_ri'):
        r += 1
        m.section(ws, r, 'Cross-check 3: residual income — CFA L2 Equity: Residual Income Valuation', 3); r += 1
        ws.cell(row=r, column=1, value='RI = NI - ke x opening book value. Terminal RI held flat in perpetuity (persistence = 1), discounted at ke.').font = F_NOTE
        r += 1
        h = r
        ws.cell(row=h, column=1, value='Year').font = F_B
        labels = ['Net income', 'Opening book value', 'Residual income', 'Discount factor (ke)', 'PV of RI']
        for j, lab in enumerate(labels):
            ws.cell(row=h + 1 + j, column=1, value=lab)
        for i, y in enumerate(YEARS):
            col = 2 + i; L = ws.cell(row=h, column=col).column_letter
            ws.cell(row=h, column=col, value=y).font = F_B
            prev = '2026E' if i == 0 else YEARS[i - 1]
            ws[f"{L}{h + 1}"] = f"={M('ni', y)}"; ws[f"{L}{h + 1}"].font = F_LINK
            ws[f"{L}{h + 2}"] = f"={M('equity', prev)}"; ws[f"{L}{h + 2}"].font = F_LINK
            ws[f"{L}{h + 3}"] = f"={L}{h + 1}-{m.ref('ke')}*{L}{h + 2}"
            ws[f"{L}{h + 4}"] = f"=1/(1+{m.ref('ke')})^{i + 1}"
            ws[f"{L}{h + 5}"] = f"={L}{h + 3}*{L}{h + 4}"
            for k in (1, 2, 3, 5):
                ws[f"{L}{h + k}"].number_format = FMT['m0']
            ws[f"{L}{h + 4}"].number_format = '0.000'
        r = h + 7
        r = m.block(ws, r, [
            ('b0', 'Book value at Dec 31, 2026E', f"={M('equity', '2026E')}", 'm0', ''),
            ('ri_pv', 'PV of 2027E–2030E residual income', f"=SUM(B{h + 5}:E{h + 5})", 'm0', ''),
            ('ri_tv', 'PV of terminal residual income', f"=E{h + 3}/{m.ref('ke')}*E{h + 4}", 'm0', ''),
            ('ri_ps', f"Residual income value per share ({c['company']['model_ccy']})", None, 'ps', 'Cross-check only'),
            ('#', 'Cross-check 4: justified forward P/E — CFA L2 Equity: Market-Based Valuation'),
            ('payout', 'Dividend payout ratio, 2027E', f"={M('dps', '2027E')}/{M('eps', '2027E')}", 'pct', ''),
            ('roe27', 'Return on equity, 2027E', f"={M('ni', '2027E')}/AVERAGE({M('equity', '2026E')},{M('equity', '2027E')})", 'pct', ''),
            ('g_sust', 'Sustainable growth: retention x ROE', None, 'pct', 'g = b x ROE'),
            ('pe_just', 'Justified forward P/E: payout / (ke - g)', None, 'x', 'n/m when g >= ke: a single-stage model cannot value this company'),
            ('pe_act', 'Actual forward P/E, 2027E', f"={m.ref('price_model')}/{M('eps', '2027E')}", 'x', ''),
        ])
        cell('ri_ps').value = f"=({m.ref('b0')}+{m.ref('ri_pv')}+{m.ref('ri_tv')})/{m.ref('shares')}"
        cell('g_sust').value = f"=(1-{m.ref('payout')})*{m.ref('roe27')}"
        cell('pe_just').value = f'=IF({m.ref("ke")}>{m.ref("g_sust")},{m.ref("payout")}/({m.ref("ke")}-{m.ref("g_sust")}),"n/m")'
    for k in ('target', 'rating', 'ret_all', 'v1', 'v2', 'blend', 'ke', 'wacc', 'dcf_ps', 'g_impl', 'price', 'mcap', 'ev'):
        m.out(k, m.ref(k))
    if c['company'].get('has_ri'):
        m.out('ri_ps', m.ref('ri_ps')); m.out('pe_just', m.ref('pe_just'))
    return ws


def sensitivity(m, arch):
    c = m.cfg
    ws = m.sheet('Sensitivity', [30] + [12] * 7)
    m.title(ws, f"{c['company']['name']} — Sensitivity", 'Every cell is a live formula on the model; nothing is pasted.')
    a = arch.sens_axes(m)
    r = 4
    m.section(ws, r, f"Price target ({c['company']['price_ccy']}, before rounding) vs. the two valuation multiples", 8); r += 1
    ws.cell(row=r, column=1, value=f"{a['m1_label']} down, {a['m2_label']} across").font = F_NOTE
    for j, v2 in enumerate(a['m2']):
        x = ws.cell(row=r, column=2 + j, value=v2); x.font = F_IN; x.number_format = a['m2_fmt']
    for i, v1 in enumerate(a['m1']):
        y = ws.cell(row=r + 1 + i, column=1, value=v1); y.font = F_IN; y.number_format = a['m1_fmt']
        for j in range(len(a['m2'])):
            m1c = f"$A{r + 1 + i}"; m2c = f"{ws.cell(row=r, column=2 + j).column_letter}${r}"
            x = ws.cell(row=r + 1 + i, column=2 + j, value=a['formula'](m1c, m2c)); x.number_format = FMT['ps']
    r += len(a['m1']) + 3
    m.section(ws, r, 'DCF value per share vs. WACC (down) and terminal growth (across)', 8); r += 1
    h = m.dcf_rows
    gs = [0.0, 0.01, 0.02, 0.03]
    ws_ = [0.06, 0.07, 0.08, 0.09, 0.10]
    for j, g in enumerate(gs):
        x = ws.cell(row=r, column=2 + j, value=g); x.font = F_IN; x.number_format = '0.0%'
    for i, w in enumerate(ws_):
        y = ws.cell(row=r + 1 + i, column=1, value=w); y.font = F_IN; y.number_format = '0.0%'
        for j in range(len(gs)):
            W = f"$A{r + 1 + i}"; G = f"{ws.cell(row=r, column=2 + j).column_letter}${r}"
            f_ = ("=(" + "+".join(f"Valuation!{col}{h + 1}/(1+{W})^{k + 1}" for k, col in enumerate('BCDE'))
                  + f"+Valuation!E{h + 1}*(1+{G})/({W}-{G})/(1+{W})^4-{m.ref('nd_ye26')}+{m.ref('inv')})/{m.ref('shares')}")
            x = ws.cell(row=r + 1 + i, column=2 + j, value=f_); x.number_format = FMT['ps']
    return ws


def ratios(m, arch):
    ws = m.sheet('Ratios', [46, 10] + [11] * 8 + [50])
    c = m.cfg
    m.title(ws, f"{c['company']['name']} — Financial analysis ratios",
            'CFA L2 Financial Statement Analysis: DuPont, cash conversion, leverage and payout. Formulas on the Model sheet.')
    r = 4
    m.header(ws, r, ['Ratio', 'Unit'] + PERIODS + ['Reading']); r += 1
    for ln in arch.ratio_lines():
        ws.cell(row=r, column=1, value=ln['label'])
        ws.cell(row=r, column=2, value=ln.get('unit', ''))
        for p in ln.get('periods', PERIODS):
            f_ = m.expand(ln['tmpl'], p, model_prefix='Model!')
            x = ws[f"{PCOL[p]}{r}"]; x.value = f'=IFERROR({f_[1:]},"")'; x.number_format = FMT[ln.get('fmt', 'pct')]
            x.border = RULE
        ws.cell(row=r, column=11, value=ln.get('note', '')).font = F_NOTE
        r += 1
    return ws


def comps(m):
    c = m.cfg
    ws = m.sheet('Comps', [22, 10, 12, 12, 12, 12, 12, 12])
    m.title(ws, f"{c['company']['name']} — Peer comparison", c['comps']['note'])
    m.header(ws, 4, c['comps']['columns'])
    for i, row in enumerate(c['comps']['rows']):
        for j, v in enumerate(row):
            x = ws.cell(row=5 + i, column=1 + j, value=v); x.font = F_B if i == 0 else F_IN
            x.border = RULE
            if j and not isinstance(v, str):
                x.number_format = '#,##0.00'
    ws.cell(row=6 + len(c['comps']['rows']), column=1, value='Source: ' + c['comps']['source']).font = F_NOTE
    return ws


def checks(m, arch):
    c = m.cfg
    ws = m.sheet('Checks', [70, 16, 16, 12])
    m.title(ws, f"{c['company']['name']} — Model integrity checks", 'Every check should read TRUE before anything is published.')
    m.header(ws, 4, ['Check', 'Model', 'Expected', 'Pass'])
    pub = c['published']
    items = [
        ('Price target equals the published target', f"={m.ref('target')}", pub['target'], 'num'),
        ('Rating equals the published rating', f"={m.ref('rating')}", pub['rating'], 'txt'),
    ] + arch.checks(m)
    r = 5
    for desc, f_, exp, kind in items:
        ws.cell(row=r, column=1, value=desc)
        a = ws.cell(row=r, column=2, value=f_); a.number_format = '#,##0.00'
        b = ws.cell(row=r, column=3, value=exp); b.font = F_IN; b.number_format = '#,##0.00'
        if kind == 'num':
            ws.cell(row=r, column=4, value=f"=ABS(B{r}-C{r})<=MAX(0.01,ABS(C{r})*0.005)")
        else:
            ws.cell(row=r, column=4, value=f"=B{r}=C{r}")
        r += 1
    ws.cell(row=r + 1, column=1, value='All checks pass').font = F_B
    ws.cell(row=r + 1, column=4, value=f"=AND(D5:D{r - 1})").font = F_B
    m.out('all_checks', f"Checks!$D${r + 1}")
    return ws


def cover(m):
    c = m.cfg; co = c['company']
    ws = m.wb.create_sheet('Cover', 0); m.sheets['Cover'] = ws
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 22; ws.column_dimensions['C'].width = 60
    ws['A1'] = 'NORTHWALL RESEARCH'; ws['A1'].font = _F(name=FONT, bold=True, size=11, color='2F5D7C')
    ws['A2'] = f"{co['name']} ({co['ticker']})"; ws['A2'].font = _F(name=FONT, bold=True, size=16, color=NAVY)
    ws['A3'] = f"{co['sector']} · Analyst: {co['analyst']} · Model version {co['version']} · Built {co['model_date']}"
    ws['A3'].font = F_NOTE
    rows = [
        ('Rating', f"={m.ref('rating')}", 'txt'),
        (f"Price target ({co['price_ccy']})", f"={m.ref('target')}", 'ps'),
        (f"Share price ({co['price_ccy']}), {c['scalars']['price']['source']}", f"={m.ref('price')}", 'ps'),
        ('Implied all-in return', f"={m.ref('ret_all')}", 'pct'),
        ('Valuation method 1', f"={m.ref('v1')}", 'ps'),
        ('Valuation method 2', f"={m.ref('v2')}", 'ps'),
        ('DCF cross-check per share', f"={m.ref('dcf_ps')}", 'ps'),
        ('Cost of equity (CAPM)', f"={m.ref('ke')}", 'pct'),
        ('Market-implied growth', f"={m.ref('g_impl')}", 'pct'),
        ('All integrity checks pass', f'=IF({m.outputs["all_checks"]},"Yes","NO - see Checks")', 'txt'),
    ]
    r = 5
    for lab, f_, fm in rows:
        ws.cell(row=r, column=1, value=lab).font = F_B
        x = ws.cell(row=r, column=2, value=f_); x.font = F_LINK; x.number_format = FMT[fm]
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='Colour code').font = F_B; r += 1
    for lab, font in [('Blue: typed input with a source (Inputs sheet only)', F_IN), ('Black: formula', F_CALC),
                      ('Green: link to another sheet', F_LINK)]:
        ws.cell(row=r, column=1, value=lab).font = font; r += 1
    ws.cell(row=r, column=1, value='Yellow fill: key assumption or output').fill = FILL_KEY; r += 2
    ws.cell(row=r, column=1, value='Sheets').font = F_B; r += 1
    for s, d in [('Inputs', 'Every typed number, with its source'), ('Model', 'Operations and financials, H1 actual + H2 estimate'),
                 ('Valuation', 'Target methods, CAPM/WACC, DCF, reverse DCF, residual income, justified P/E'),
                 ('NAV', 'Asset-life cash flows (resource names)'), ('Sensitivity', 'Target and DCF grids'),
                 ('Ratios', 'DuPont, cash conversion, leverage, payout'), ('Comps', 'Peer table'),
                 ('Checks', 'Integrity checks; all must be TRUE'), ('Log', 'Every change, dated, with its reason')]:
        if s in m.wb.sheetnames or s in ('NAV',):
            ws.cell(row=r, column=1, value=s); ws.cell(row=r, column=2, value=d).font = F_NOTE; r += 1


def log(m):
    ws = m.sheet('Log', [12, 14, 60, 30])
    m.title(ws, 'Model change log', 'Newest first. One row per change; never delete rows.')
    m.header(ws, 4, ['Date', 'Version', 'Change', 'Effect on target'])
    for i, e in enumerate(m.cfg['log']):
        for j, k in enumerate(['date', 'version', 'change', 'effect']):
            ws.cell(row=5 + i, column=1 + j, value=e[k]).font = F_IN
