"""Shared asset-life NAV sheet writer (CFA L2: valuation of resource companies is an
application of DCF over a finite asset life; see docs/CFA_Modelling_Framework.md)."""
from openpyxl.utils import get_column_letter
from engine import FMT, F_B, F_LINK, F_NOTE, RULE


def nav_sheet(m, title, sub, years, rows, first_col=3):
    ws = m.sheet('NAV', [34, 10] + [9] * len(years))
    m.title(ws, title, sub)
    r0 = 4
    ws.cell(row=r0, column=1, value='Year').font = F_B
    col = {}
    for i, y in enumerate(years):
        L = get_column_letter(first_col + i); col[y] = L
        ws[f"{L}{r0}"] = y; ws[f"{L}{r0}"].font = F_B
    rowno = {}
    for j, rw in enumerate(rows):
        rowno[rw[0]] = r0 + 1 + j
    for j, (key, label, unit, fmt, fn) in enumerate(rows):
        rr = r0 + 1 + j
        ws.cell(row=rr, column=1, value=label); ws.cell(row=rr, column=2, value=unit)
        for i, y in enumerate(years):
            L = col[y]; P = col.get(years[i - 1]) if i else None
            f = fn(y, L, P, rowno)
            c = ws[f"{L}{rr}"]; c.value = f; c.number_format = FMT[fmt]; c.border = RULE
            if isinstance(f, str) and f.startswith('=Model!'):
                c.font = F_LINK
    first, last = col[years[0]], col[years[-1]]
    return ws, rowno, r0 + len(rows) + 2, first, last
