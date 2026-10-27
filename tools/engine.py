"""Northwall model engine.

Turns a company's documented inputs (coverage/<TICKER>/model_inputs.yaml) plus the
house price deck (data/price_deck.yaml) into a formula-driven Excel model laid out
the way a sell-side model is: Cover, Inputs, Model, Valuation, NAV (where relevant),
Sensitivity, Ratios, Comps, Checks, Log.

Conventions (CFA / bank style):
  blue font  = typed input (every one carries a source in the Inputs sheet)
  black font = formula
  green font = link to another sheet
No number in Model, Valuation or Sensitivity is typed; they all trace to Inputs.
"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

PERIODS = ['2025A', 'H1-26A', 'H2-26E', '2026E', '2027E', '2028E', '2029E', '2030E']
PTYPE = {'2025A': 'A', 'H1-26A': 'H1', 'H2-26E': 'H2', '2026E': 'FY',
         '2027E': 'F', '2028E': 'F', '2029E': 'F', '2030E': 'F'}
PCOL = {p: get_column_letter(3 + i) for i, p in enumerate(PERIODS)}          # C..J
PREV = {'H1-26A': '2025A', 'H2-26E': 'H1-26A', '2026E': '2025A', '2027E': '2026E',
        '2028E': '2027E', '2029E': '2028E', '2030E': '2029E'}
FORECAST = ['2027E', '2028E', '2029E', '2030E']

NAVY = '0F1A2B'; STEEL = '2F5D7C'; SLATE = '5A6B7C'
FONT = 'Arial'
F_IN = Font(name=FONT, color='0000FF')
F_CALC = Font(name=FONT, color='000000')
F_LINK = Font(name=FONT, color='008000')
F_B = Font(name=FONT, bold=True)
F_HEAD = Font(name=FONT, bold=True, color='F7F8FA')
F_NOTE = Font(name=FONT, italic=True, size=8, color=SLATE)
FILL_HEAD = PatternFill('solid', fgColor=NAVY)
FILL_SEC = PatternFill('solid', fgColor='E9EEF2')
FILL_KEY = PatternFill('solid', fgColor='FFFF00')
RULE = Border(bottom=Side(style='thin', color='D9DEE3'))
FMT = {'m': '#,##0.0;(#,##0.0);-', 'm0': '#,##0;(#,##0);-', 'ps': '#,##0.00;(#,##0.00);-',
       'pct': '0.0%;(0.0%);-', 'x': '0.0x', 'px': '#,##0;(#,##0);-', 'n1': '#,##0.0;(#,##0.0);-',
       'n2': '#,##0.00;(#,##0.00);-', 'n3': '0.000', 'int': '0', 'txt': '@'}


def is_formula(v):
    return isinstance(v, str) and v.startswith('=')


class Model:
    def __init__(self, cfg, deck):
        self.cfg = cfg
        self.deck = deck
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sheets = {}
        self.series = {}     # key -> row on Inputs (period series)
        self.has = {}        # key -> set of periods with a value
        self.scal = {}       # key -> row on Inputs (scalars)
        self.rows = {}       # key -> row on Model
        self.cells = {}      # name -> 'Sheet!$C$5' absolute refs for named outputs
        self.outputs = {}    # name -> absolute ref read back after recalc

    # ------------------------------------------------------------------ helpers
    def sheet(self, name, widths=None):
        ws = self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths or []):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        self.sheets[name] = ws
        return ws

    def title(self, ws, text, sub):
        ws['A1'] = text; ws['A1'].font = Font(name=FONT, bold=True, size=14, color=NAVY)
        ws['A2'] = sub; ws['A2'].font = F_NOTE

    def header(self, ws, row, labels, start_col=1):
        for i, t in enumerate(labels):
            c = ws.cell(row=row, column=start_col + i, value=t)
            c.font = F_HEAD; c.fill = FILL_HEAD
            c.alignment = Alignment(horizontal='center' if i else 'left', vertical='center', wrap_text=True)

    def section(self, ws, row, text, ncols=11):
        for col in range(1, ncols + 1):
            ws.cell(row=row, column=col).fill = FILL_SEC
        c = ws.cell(row=row, column=1, value=text); c.font = Font(name=FONT, bold=True, color=NAVY)

    def S(self, key):           # absolute ref to a scalar input
        return f"Inputs!$C${self.scal[key]}"

    def I(self, key, period):   # ref to a period input
        return f"Inputs!{PCOL[period]}{self.series[key]}"

    def M(self, key, period):   # absolute ref to a model line
        return f"Model!${PCOL[period]}${self.rows[key]}"

    def expand(self, tmpl, period, model_prefix=''):
        """Template tokens: [k] model line same period, [k<] model line previous period,
        {k} period input same period, {{k}} scalar input."""
        def rep_s(m): return self.S(m.group(1))
        def rep_i(m): return self.I(m.group(1), period)
        def rep_prev(m): return f"{model_prefix}{PCOL[PREV[period]]}{self.rows[m.group(1)]}"
        def rep_m(m): return f"{model_prefix}{PCOL[period]}{self.rows[m.group(1)]}"
        t = re.sub(r'\{\{(\w+)\}\}', rep_s, tmpl)
        t = re.sub(r'\{(\w+)\}', rep_i, t)
        t = re.sub(r'\[(\w+)<\]', rep_prev, t)
        t = re.sub(r'\[(\w+)\]', rep_m, t)
        return t

    # ------------------------------------------------------------------ Inputs
    def build_inputs(self):
        ws = self.sheet('Inputs', [38, 10] + [11] * 8 + [70])
        c = self.cfg
        self.title(ws, f"{c['company']['name']} ({c['company']['ticker']}) — Inputs",
                   'Every typed number in the model lives here, in blue, with its source. Change inputs here only.')
        r = 4
        self.section(ws, r, 'House price deck (from data/price_deck.yaml)'); r += 1
        self.header(ws, r, ['Input', 'Unit'] + PERIODS + ['Source']); r += 1
        for key, d in self.deck['series'].items():
            if key not in c.get('deck_uses', []):
                continue
            ws.cell(row=r, column=1, value=d['label']).font = F_B
            ws.cell(row=r, column=2, value=d['unit'])
            for p in PERIODS:
                v = d['values'].get(p)
                if v is None:
                    continue
                cell = ws[f"{PCOL[p]}{r}"]; cell.value = v; cell.font = F_IN
                cell.number_format = FMT.get(d.get('fmt', 'm0'))
            ws.cell(row=r, column=11, value=d['source']).font = F_NOTE
            self.series[key] = r; self.has[key] = set(d['values']); r += 1
        r += 1
        self.section(ws, r, 'Company drivers by period'); r += 1
        self.header(ws, r, ['Input', 'Unit'] + PERIODS + ['Source']); r += 1
        for key, d in c['series'].items():
            ws.cell(row=r, column=1, value=d['label'])
            ws.cell(row=r, column=2, value=d.get('unit', ''))
            for p, v in d['values'].items():
                cell = ws[f"{PCOL[p]}{r}"]
                cell.value = v
                cell.font = F_CALC if is_formula(v) else F_IN
                cell.number_format = FMT.get(d.get('fmt', 'm'))
            ws.cell(row=r, column=11, value=d['source']).font = F_NOTE
            self.series[key] = r; self.has[key] = set(d['values']); r += 1
        r += 1
        self.section(ws, r, 'Scalar inputs: market data, operating assumptions, valuation parameters'); r += 1
        self.header(ws, r, ['Input', 'Unit', 'Value'] + [''] * 7 + ['Source']); r += 1
        for key, d in c['scalars'].items():
            ws.cell(row=r, column=1, value=d['label'])
            ws.cell(row=r, column=2, value=d.get('unit', ''))
            cell = ws[f"C{r}"]; cell.value = d['value']
            cell.font = F_CALC if is_formula(d['value']) else F_IN
            cell.number_format = FMT.get(d.get('fmt', 'n2'))
            if d.get('key_assumption'):
                cell.fill = FILL_KEY
            ws.cell(row=r, column=11, value=d['source']).font = F_NOTE
            self.scal[key] = r; r += 1
        # deck scalars (rf, erp, etc.)
        for key, d in self.deck.get('scalars', {}).items():
            if key not in c.get('deck_uses', []):
                continue
            ws.cell(row=r, column=1, value=d['label'])
            ws.cell(row=r, column=2, value=d.get('unit', ''))
            cell = ws[f"C{r}"]; cell.value = d['value']; cell.font = F_IN
            cell.number_format = FMT.get(d.get('fmt', 'pct'))
            ws.cell(row=r, column=11, value=d['source']).font = F_NOTE
            self.scal[key] = r; r += 1
        ws.freeze_panes = 'C4'

    # ------------------------------------------------------------------ Model
    def build_model(self, lines, ccy):
        ws = self.sheet('Model', [40, 10] + [11] * 8 + [52])
        c = self.cfg['company']
        self.title(ws, f"{c['name']} — Operating and financial model ({ccy} millions unless noted)",
                   'H1-26A = reported six months; H2-26E = Northwall estimate; 2026E = H1 + H2. All cells are formulas.')
        r = 4
        self.header(ws, r, ['Line', 'Unit'] + PERIODS + ['Method']); r += 1
        # first pass: assign rows so templates can reference lines defined later
        rr = r
        for ln in lines:
            if 'key' not in ln:
                rr += 1; continue
            self.rows[ln['key']] = rr; rr += 1
        for ln in lines:
            if 'key' not in ln:
                self.section(ws, r, ln['section']); r += 1
                continue
            ws.cell(row=r, column=1, value=ln['label']).font = F_B if ln.get('bold') else F_CALC
            ws.cell(row=r, column=2, value=ln.get('unit', ''))
            spec = ln['spec']
            for p in PERIODS:
                t = PTYPE[p]
                tmpl = spec.get(p, spec.get(t, spec.get('all')))
                if tmpl is None:
                    continue
                if tmpl == 'sum':
                    f = f"={PCOL['H1-26A']}{r}+{PCOL['H2-26E']}{r}"
                elif tmpl == 'last':
                    f = f"={PCOL['H2-26E']}{r}"
                elif tmpl.startswith('wavg:'):
                    w = self.rows[tmpl[5:]]; h1, h2, fy = PCOL['H1-26A'], PCOL['H2-26E'], PCOL['2026E']
                    f = f"=({h1}{r}*{h1}{w}+{h2}{r}*{h2}{w})/{fy}{w}"
                else:
                    missing = [k for k in re.findall(r'(?<!\{)\{(\w+)\}(?!\})', tmpl) if p not in self.has.get(k, set())]
                    if missing:
                        continue                       # no input for this period: leave blank
                    f = self.expand(tmpl, p)
                cell = ws[f"{PCOL[p]}{r}"]; cell.value = f
                cell.font = F_LINK if re.fullmatch(r"=Inputs![A-Z]+\d+", f) else F_CALC
                if ln.get('bold'):
                    cell.font = Font(name=FONT, bold=True)
                cell.number_format = FMT.get(ln.get('fmt', 'm'))
                cell.border = RULE
            ws.cell(row=r, column=11, value=ln.get('note', '')).font = F_NOTE
            r += 1
        ws.freeze_panes = 'C5'

    # ------------------------------------------------------------------ generic block writer
    def block(self, ws, r, items, label_w=None):
        """items: (key, label, value_or_formula, fmt, note). Writes label in A, value in B, note in C."""
        for it in items:
            if it[0] == '#':
                self.section(ws, r, it[1], 3); r += 1
                continue
            key, label, v, fmt, note = it
            ws.cell(row=r, column=1, value=label)
            cell = ws.cell(row=r, column=2, value=v)
            cell.font = F_CALC if is_formula(v) else F_IN
            if isinstance(v, str) and re.fullmatch(r"=(Inputs|Model)!\$?[A-Z]+\$?\d+", v or ''):
                cell.font = F_LINK
            cell.number_format = FMT.get(fmt, fmt)
            cell.border = RULE
            ws.cell(row=r, column=3, value=note).font = F_NOTE
            if key:
                self.cells[key] = f"'{ws.title}'!$B${r}"
            r += 1
        return r

    def ref(self, key):
        return self.cells[key]

    def out(self, name, ref):
        self.outputs[name] = ref

    def save(self, path):
        from copy import copy
        for ws in self.wb.worksheets:                 # one professional font throughout
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None and (c.font is None or c.font.name != FONT):
                        f = copy(c.font); f.name = FONT; c.font = f
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(path)
